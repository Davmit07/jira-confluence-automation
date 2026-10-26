"""Generate synthetic Vehicle Shop draft cases, without fetching Jira stories."""

import argparse
import os
import tempfile
from pathlib import Path
from typing import Union
from urllib.parse import urlsplit

from openpyxl import load_workbook

from create_test_catalog import HEADERS, create_test_catalog
from validate_test_catalog import validate_configuration, validate_test_catalog


DEFAULT_OUTPUT = Path("reports/vehicle-shop-test-catalog.xlsx")
DEFAULT_JIRA_URL = "https://jiraeu.epam.com/projects/EPMCDMETST/"
DEFAULT_PROJECT_KEY = "EPMCDMETST"

# These issue numbers are illustrative, not retrieved or verified Jira stories.
SCENARIOS = [
    ("Search", "High", "Search available vehicles", "vehicles.available",
     "1. Open vehicle search\n2. Select the synthetic make and model\n3. Submit the search",
     "Only matching available vehicles appear, with consistent vehicle details."),
    ("Search", "Medium", "Search with no matching vehicles", "vehicles.no-match",
     "1. Open vehicle search\n2. Select criteria with no seeded matches\n3. Submit the search\n4. Clear the criteria",
     "An empty-state message appears; clearing criteria restores available vehicles."),
    ("Search", "High", "Recover from a vehicle search UI failure", "vehicles.search-failure",
     "1. Enable the approved search-failure UI fixture\n2. Open vehicle search\n3. Submit valid criteria\n4. Reset the fixture and retry",
     "An actionable error appears without stale results; retry recovers after reset."),
    ("Compare", "High", "Compare two available vehicles", "vehicles.compare-pair",
     "1. Open vehicle search\n2. Select two distinct synthetic vehicles for comparison\n3. Open comparison\n4. Remove one vehicle",
     "Both vehicles show corresponding specifications; removal updates the comparison."),
    ("Compare", "Medium", "Prevent duplicate comparison selection", "vehicles.compare-pair",
     "1. Open vehicle search\n2. Add a synthetic vehicle to comparison\n3. Attempt to add the same vehicle again",
     "The comparison contains one entry for that vehicle and gives clear selection feedback."),
    ("Compare", "High", "Handle an unavailable comparison vehicle", "vehicles.compare-unavailable",
     "1. Seed a comparison with an unavailable synthetic vehicle\n2. Open comparison\n3. Remove the unavailable vehicle\n4. Add an available replacement",
     "Unavailable data is identified without misleading specifications; replacement recovers the comparison."),
    ("Configure", "High", "Configure compatible vehicle options", "vehicle.compatible-options",
     "1. Open a synthetic vehicle configurator\n2. Choose compatible trim and options\n3. Review the summary\n4. Reset the configuration",
     "Selected options and displayed total are consistent; reset restores the initial state."),
    ("Configure", "High", "Reject incompatible configuration options", "vehicle.incompatible-options",
     "1. Open a synthetic vehicle configurator\n2. Select the first option in a known incompatible pair\n3. Attempt to select the second option",
     "The UI explains the conflict and prevents an invalid combination and misleading total."),
    ("Configure", "High", "Recover from unavailable configuration pricing", "vehicle.pricing-failure",
     "1. Enable the approved pricing-failure UI fixture\n2. Open a synthetic vehicle configurator\n3. Change an option\n4. Reset the fixture and retry",
     "Pricing failure is visible; no stale price is presented as confirmed and retry restores pricing."),
    ("Booking", "Critical", "Create a synthetic Sandbox booking", "booking.synthetic",
     "1. Open the approved Sandbox vehicle booking page\n2. Select a seeded vehicle and test slot\n3. Enter named synthetic contact data\n4. Submit through the approved simulated booking flow\n5. Remove the test booking using approved cleanup",
     "A Sandbox-only simulated confirmation appears; no live payment occurs and cleanup restores the slot."),
    ("Booking", "High", "Require mandatory booking fields", "booking.synthetic",
     "1. Open the approved Sandbox vehicle booking page\n2. Select a seeded vehicle and test slot\n3. Leave required contact fields empty\n4. Attempt to continue",
     "Field-level validation identifies missing values; no booking or payment is created."),
    ("Booking", "High", "Handle a booking slot becoming unavailable", "booking.unavailable-slot",
     "1. Open the approved Sandbox vehicle booking page\n2. Select a synthetic slot\n3. Enable the approved slot-unavailable UI fixture\n4. Attempt the simulated booking\n5. Reset the fixture",
     "The UI explains slot unavailability without confirmation, duplicate booking, or payment."),
]


def create_shop_catalog(
    output_path: Union[str, Path] = DEFAULT_OUTPUT,
    jira_base_url: str = DEFAULT_JIRA_URL,
    project_key: str = DEFAULT_PROJECT_KEY,
    overwrite: bool = False,
) -> Path:
    """Publish a validated sample atomically; never overwrite without consent."""
    validate_configuration(jira_base_url, project_key)
    output_path = Path(output_path)
    if output_path.suffix.lower() != ".xlsx":
        raise ValueError("Output must have an .xlsx extension.")
    if output_path.exists() and not overwrite:
        raise FileExistsError("Catalog already exists; choose another path or pass --force.")
    if not output_path.parent.is_dir():
        raise FileNotFoundError("Output directory does not exist.")
    parsed = urlsplit(jira_base_url)
    jira_origin_url = f"https://{parsed.netloc}"
    with tempfile.TemporaryDirectory(prefix=".shop-catalog-", dir=output_path.parent) as directory:
        temporary_path = Path(directory) / "catalog.xlsx"
        create_test_catalog(temporary_path)
        workbook = load_workbook(temporary_path)
        try:
            sheet = workbook["Test Cases"]
            for number, (journey, priority, title, data, steps, expected) in enumerate(SCENARIOS, start=1):
                issue_number = (number - 1) // 3 + 1
                issue_key = f"{project_key}-{issue_number}"
                preconditions = (
                    "SYNTHETIC EXAMPLE; Jira story and acceptance criteria are not verified. "
                    "QA must confirm expected behavior, approved UI fixtures, test data, and cleanup. "
                    "Use an allowlisted non-production environment and dedicated test account."
                )
                if journey == "Booking":
                    preconditions += (
                        " Sandbox only; application-owner approval and simulated booking/payment "
                        "are required before execution. Never use real customer or payment data."
                    )
                values = {
                    "Case ID": f"SHOP-{number:03d}",
                    "Jira Issue Key": issue_key,
                    "Jira Issue Link": f"{jira_origin_url}/browse/{issue_key}",
                    "Journey": journey,
                    "Title": f"[SYNTHETIC] {title}",
                    "Priority": priority,
                    "Preconditions": preconditions,
                    "Steps": steps,
                    "Expected Results": expected,
                    "Test Data Reference": f"synthetic:{data}",
                    "Review Status": "Draft",
                    "Automation Status": "Not Mapped",
                    "Last Result": "Not Run",
                }
                for column, header in enumerate(HEADERS, start=1):
                    sheet.cell(row=number + 1, column=column, value=values.get(header))
                sheet.row_dimensions[number + 1].height = 150
            workbook["Read Me"].append((
                "Synthetic examples",
                "All cases and issue numbers are illustrative, not fetched or verified Jira "
                "requirements. Replace references and confirm acceptance criteria before QA "
                "approval. Named data/fixtures are proposals, not implemented automation.",
            ))
            workbook["Read Me"].append((
                "Execution gates",
                "Draft / Not Mapped cases cannot run. Validation is not approval or proof of "
                "implemented automation. The future runner must enforce a fixed automation "
                "registry, explicit confirmation, production denial, and exact-origin "
                "allowlisting before navigation, including Sandbox-only bookings.",
            ))
            workbook.properties.title = "Vehicle Shop Synthetic Draft Test Catalog"
            workbook.save(temporary_path)
        finally:
            workbook.close()
        validation = validate_test_catalog(temporary_path, jira_base_url, project_key)
        if validation["outcome"] != "valid":
            raise ValueError(f"Generated catalog failed validation: {validation['errors']}")
        if overwrite:
            os.replace(temporary_path, output_path)
        else:
            # A hard link publishes only if the destination still does not exist.
            os.link(temporary_path, output_path)
    return output_path


def main() -> None:
    parser = argparse.ArgumentParser(description="Create a validated synthetic shop catalog; no Jira access.")
    parser.add_argument("output", nargs="?", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--jira-base-url", default=DEFAULT_JIRA_URL)
    parser.add_argument("--project-key", default=DEFAULT_PROJECT_KEY)
    parser.add_argument("--force", action="store_true")
    arguments = parser.parse_args()
    try:
        output = create_shop_catalog(
            arguments.output, arguments.jira_base_url, arguments.project_key, arguments.force
        )
    except (OSError, ValueError) as error:
        parser.error(str(error))
    print(f"Created validated synthetic draft catalog ({len(SCENARIOS)} cases): {output}")


if __name__ == "__main__":
    main()
