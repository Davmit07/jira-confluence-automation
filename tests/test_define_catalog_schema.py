import runpy
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from catalog_schema import DEFAULTS, FIELD_RULES, HEADERS, LISTS, REQUIRED_FIELDS, SCHEMA_VERSION
from validate_test_catalog import validate_test_catalog


SCRIPT = Path(__file__).resolve().parents[3] / "tools" / "define-corepy"
define_catalog = runpy.run_path(str(SCRIPT))["define_catalog"]


class DefineCatalogSchemaTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.output = Path(self.directory.name) / "schema.xlsx"

    def test_shared_schema_contract(self):
        self.assertEqual(SCHEMA_VERSION, "1.0.0")
        self.assertEqual(set(FIELD_RULES), set(HEADERS))
        self.assertTrue(set(REQUIRED_FIELDS).issubset(HEADERS))
        self.assertTrue(set(LISTS).issubset(HEADERS))
        self.assertEqual(len(HEADERS), 16)
        self.assertEqual(len(set(HEADERS)), 16)

    def test_creates_versioned_schema_and_incomplete_drafts(self):
        define_catalog(["SHOP-123", "SHOP-456"], self.output, "https://jira.example:8443/projects/SHOP/")
        workbook = load_workbook(self.output)
        try:
            self.assertEqual(workbook["Schema"]["B1"].value, SCHEMA_VERSION)
            self.assertEqual(
                [row[0] for row in workbook["Schema"].iter_rows(min_row=3, values_only=True)],
                HEADERS,
            )
            self.assertIn(("Schema Version", SCHEMA_VERSION), list(workbook["Read Me"].values))
            sheet = workbook["Test Cases"]
            self.assertEqual([cell.value for cell in sheet[1]], HEADERS)
            for index, key in enumerate(("SHOP-123", "SHOP-456"), start=2):
                row = dict(zip(HEADERS, [cell.value for cell in sheet[index]]))
                self.assertEqual(row["Case ID"], f"{key}-TC-001")
                self.assertEqual(row["Jira Issue Key"], key)
                self.assertEqual(row["Jira Issue Link"], f"https://jira.example:8443/browse/{key}")
                for header, default in DEFAULTS.items():
                    self.assertEqual(row[header], default)
                self.assertIsNone(row["Title"])
                self.assertIsNone(row["Automation ID"])
                self.assertIsNone(row["Reviewer"])
            self.assertEqual(len(sheet.data_validations.dataValidation), len(LISTS))
        finally:
            workbook.close()
        result = validate_test_catalog(self.output, "https://jira.example:8443", "SHOP")
        self.assertEqual(result["outcome"], "invalid")
        self.assertEqual(result["checked_rows"], 2)
        self.assertEqual(
            {error["column"] for error in result["errors"]},
            {"Journey", "Title", "Priority", "Preconditions", "Steps", "Expected Results"},
        )

    def test_invalid_keys_and_configuration_never_write(self):
        for keys in ([], ["shop-1"], ["SHOP-0"], ["SHOP-01"], ["SHOP-1", "SHOP-1"],
                     ["SHOP-1", "OTHER-2"], ["=1+1"], ["SHOP-1 "], ["SHOP-1"] * 1001):
            with self.subTest(keys=keys[:2]):
                with self.assertRaises(ValueError):
                    define_catalog(keys, self.output)
                self.assertFalse(self.output.exists())
        with self.assertRaises(ValueError):
            define_catalog(["SHOP-1"], self.output, "http://jira.example")
        with self.assertRaises(ValueError):
            define_catalog(["SHOP-1"], self.output.with_suffix(".xls"))
        with self.assertRaises(FileNotFoundError):
            define_catalog(["SHOP-1"], self.output.parent / "missing" / "schema.xlsx")

    def test_overwrite_requires_force(self):
        self.output.write_bytes(b"user workbook")
        with self.assertRaises(FileExistsError):
            define_catalog(["SHOP-1"], self.output)
        self.assertEqual(self.output.read_bytes(), b"user workbook")
        define_catalog(["SHOP-1"], self.output, overwrite=True)
        workbook = load_workbook(self.output)
        self.assertEqual(workbook["Schema"]["B1"].value, SCHEMA_VERSION)
        workbook.close()

    def test_failed_save_preserves_existing_file_and_cleans_temporary_directory(self):
        self.output.write_bytes(b"user workbook")
        with patch("openpyxl.workbook.workbook.Workbook.save", side_effect=OSError("write failed")):
            with self.assertRaisesRegex(OSError, "write failed"):
                define_catalog(["SHOP-1"], self.output, overwrite=True)
        self.assertEqual(self.output.read_bytes(), b"user workbook")
        self.assertEqual(list(self.output.parent.iterdir()), [self.output])

    def test_cli_from_another_working_directory_and_explicit_errors(self):
        result = subprocess.run(
            [sys.executable, str(SCRIPT), "SHOP-123", "SHOP-456",
             "--output", str(self.output), "--jira-base-url", "https://jira.example"],
            cwd=self.directory.name, capture_output=True, text=True,
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("v1.0.0", result.stdout)
        self.assertIn("required story fields remain incomplete", result.stdout)
        for arguments in ([], ["bad-key"], ["SHOP-1", "--output", str(self.output)]):
            result = subprocess.run(
                [sys.executable, str(SCRIPT), *arguments],
                cwd=self.directory.name, capture_output=True, text=True,
            )
            self.assertEqual(result.returncode, 2)
            self.assertIn("error:", result.stderr)
            self.assertNotIn("Traceback", result.stderr)


if __name__ == "__main__":
    unittest.main()
