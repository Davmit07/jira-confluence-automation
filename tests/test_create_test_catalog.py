import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from create_test_catalog import HEADERS, LISTS, create_test_catalog


class CreateTestCatalogTests(unittest.TestCase):
    def test_creates_catalog_with_traceability_columns_and_dropdowns(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            output_path = Path(temporary_directory) / "catalog.xlsx"
            create_test_catalog(output_path)

            workbook = load_workbook(output_path)
            self.assertEqual(workbook.sheetnames, ["Test Cases", "Read Me"])
            cases = workbook["Test Cases"]
            self.assertEqual(
                [
                    cases.cell(row=1, column=column).value
                    for column in range(1, len(HEADERS) + 1)
                ],
                HEADERS,
            )
            self.assertEqual(cases.freeze_panes, "A2")
            self.assertIsNone(cases["A2"].value)
            self.assertIn("Review Date", HEADERS)

            validations = {
                validation.formula1.strip('"')
                for validation in cases.data_validations.dataValidation
            }
            for options in LISTS.values():
                self.assertIn(",".join(options), validations)
            self.assertEqual(cases.auto_filter.ref, "A1:P1001")

    def test_refuses_to_overwrite_without_force(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            output_path = Path(temporary_directory) / "catalog.xlsx"
            output_path.write_text("existing", encoding="utf-8")

            with self.assertRaises(FileExistsError):
                create_test_catalog(output_path)

            self.assertEqual(output_path.read_text(encoding="utf-8"), "existing")


if __name__ == "__main__":
    unittest.main()
