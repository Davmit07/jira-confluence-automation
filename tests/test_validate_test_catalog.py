import contextlib
import hashlib
import io
import json
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from create_test_catalog import HEADERS, LISTS
from validate_test_catalog import REQUIRED_FIELDS, main, validate_test_catalog


JIRA_URL = "https://jiraeu.epam.com/projects/EPMCDMETST/"
PROJECT = "EPMCDMETST"


def valid_case(**changes):
    values = dict(zip(HEADERS, [
        "SHOP-001", "EPMCDMETST-1", "https://jiraeu.epam.com/browse/EPMCDMETST-1",
        "Search", "Find a synthetic vehicle", "High", "Synthetic inventory is seeded.",
        "1. Open search\n2. Submit criteria", "Matching vehicles appear.", None,
        "Draft", None, None, "Not Mapped", None, "Not Run",
    ]))
    values.update(changes)
    return values


class ValidateTestCatalogTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name) / "catalog.xlsx"

    def write(self, rows=None, headers=None):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Test Cases"
        headers = HEADERS if headers is None else headers
        sheet.append(headers)
        for row in [valid_case()] if rows is None else rows:
            sheet.append([row.get(header) for header in headers])
        workbook.save(self.path)
        workbook.close()

    def validate(self):
        return validate_test_catalog(self.path, JIRA_URL, PROJECT)

    def assert_column_error(self, result, column, row=2):
        self.assertEqual(result["outcome"], "invalid")
        self.assertTrue(any(
            error["column"] == column and error["row"] == row
            for error in result["errors"]
        ), result)
        self.assertEqual(result["error_count"], len(result["errors"]))

    def test_valid_result_shape_and_read_only(self):
        self.write()
        before = hashlib.sha256(self.path.read_bytes()).hexdigest()
        self.assertEqual(self.validate(), {
            "outcome": "valid", "checked_rows": 1, "error_count": 0, "errors": [],
        })
        self.assertEqual(hashlib.sha256(self.path.read_bytes()).hexdigest(), before)

    def test_invalid_validation_is_also_read_only(self):
        self.write([valid_case(Title="")])
        before = self.path.read_bytes()
        self.assert_column_error(self.validate(), "Title")
        self.assertEqual(self.path.read_bytes(), before)

    def test_headers_can_be_reordered_and_blank_rows_are_ignored(self):
        self.write([{}, valid_case(), {"Title": "  "}, valid_case(**{"Case ID": "SHOP-002"})],
                   list(reversed(HEADERS)))
        self.assertEqual(self.validate()["checked_rows"], 2)
        self.assertEqual(self.validate()["outcome"], "valid")

    def test_empty_template_is_valid_but_checks_zero_cases(self):
        self.write([])
        self.assertEqual(self.validate()["checked_rows"], 0)
        self.assertEqual(self.validate()["outcome"], "valid")

    def test_every_required_field(self):
        for header in REQUIRED_FIELDS:
            with self.subTest(header=header):
                self.write([valid_case(**{header: " "})])
                self.assert_column_error(self.validate(), header)

    def test_optional_fields_remain_optional(self):
        self.write()
        self.assertEqual(self.validate()["outcome"], "valid")

    def test_case_ids_are_trimmed_and_case_insensitive_for_uniqueness(self):
        self.write([valid_case(), {}, valid_case(**{"Case ID": " shop-001 "})])
        result = self.validate()
        self.assert_column_error(result, "Case ID", row=4)
        self.assertIn("row 2", result["errors"][0]["message"])

    def test_all_enum_values_and_rejection_without_normalization(self):
        for header, options in LISTS.items():
            for value in options:
                with self.subTest(header=header, value=value):
                    self.write([valid_case(**{
                        header: value, "Reviewer": "QA owner",
                        "Review Date": date(2026, 10, 26), "Automation ID": "search.example",
                    })])
                    self.assertEqual(self.validate()["outcome"], "valid")
            for value in ("unknown", options[0].lower(), options[0] + " "):
                with self.subTest(header=header, invalid=value):
                    self.write([valid_case(**{header: value})])
                    self.assert_column_error(self.validate(), header)

    def test_invalid_issue_keys(self):
        for key in ("OTHER-1", "epmcdmetst-1", "EPMCDMETST-0", "EPMCDMETST-01",
                    "EPMCDMETST-a", " EPMCDMETST-1", "EPMCDMETST-1 "):
            with self.subTest(key=key):
                self.write([valid_case(**{"Jira Issue Key": key})])
                self.assert_column_error(self.validate(), "Jira Issue Key")

    def test_invalid_jira_links(self):
        for link in (
            "/browse/EPMCDMETST-1",
            "http://jiraeu.epam.com/browse/EPMCDMETST-1",
            "https://jiraeu.epam.com.evil.example/browse/EPMCDMETST-1",
            "https://jiraeu.epam.com:444/browse/EPMCDMETST-1",
            "https://user:secret@jiraeu.epam.com/browse/EPMCDMETST-1",
            "https://jiraeu.epam.com/browse/EPMCDMETST-2",
            "https://jiraeu.epam.com/projects/EPMCDMETST/",
            "https://jiraeu.epam.com/browse/EPMCDMETST-1?token=secret",
            "https://jiraeu.epam.com/browse/EPMCDMETST-1#fragment",
            "https://jiraeu.epam.com:bad/browse/EPMCDMETST-1",
            "https://jiraeu.epam.com\\evil/browse/EPMCDMETST-1",
            " https://jiraeu.epam.com/browse/EPMCDMETST-1",
        ):
            with self.subTest(link=link):
                self.write([valid_case(**{"Jira Issue Link": link})])
                result = self.validate()
                self.assert_column_error(result, "Jira Issue Link")
                self.assertNotIn("secret", json.dumps(result))

    def test_explicit_https_default_port_and_hostname_case_are_same_origin(self):
        self.write([valid_case(**{
            "Jira Issue Link": "https://JIRAEU.EPAM.COM:443/browse/EPMCDMETST-1"
        })])
        self.assertEqual(self.validate()["outcome"], "valid")

    def test_numbered_steps(self):
        for steps in (
            "Open search", "2. Open search", "1. Open\n1. Submit",
            "1. Open\n3. Submit", "2. Submit\n1. Open", "1. ",
            "1. Open\n\n2. Submit", "1. Open\nun-numbered continuation",
        ):
            with self.subTest(steps=steps):
                self.write([valid_case(Steps=steps)])
                self.assert_column_error(self.validate(), "Steps")
        self.write([valid_case(Steps="1) Open search\n2) Submit")])
        self.assertEqual(self.validate()["outcome"], "valid")

    def test_approval_requires_reviewer_and_date(self):
        self.write([valid_case(**{"Review Status": "Approved"})])
        result = self.validate()
        self.assert_column_error(result, "Reviewer")
        self.assert_column_error(result, "Review Date")

    def test_review_dates(self):
        for value in ("2026-02-30", "26/10/2026", 123, " "):
            with self.subTest(value=value):
                self.write([valid_case(**{"Review Date": value})])
                self.assert_column_error(self.validate(), "Review Date")
        self.write([valid_case(**{
            "Review Status": "Approved", "Reviewer": "QA owner", "Review Date": "2026-10-26",
        })])
        self.assertEqual(self.validate()["outcome"], "valid")

    def test_implemented_requires_an_id_but_does_not_claim_registry_validation(self):
        self.write([valid_case(**{"Automation Status": "Implemented"})])
        self.assert_column_error(self.validate(), "Automation ID")

    def test_multiple_errors_across_rows_are_collected(self):
        self.write([valid_case(Title="", Steps="bad"), valid_case(Priority="urgent")])
        result = self.validate()
        self.assertEqual(result["checked_rows"], 2)
        self.assertEqual(result["error_count"], 4)
        self.assert_column_error(result, "Title")
        self.assert_column_error(result, "Steps")
        self.assert_column_error(result, "Case ID", 3)
        self.assert_column_error(result, "Priority", 3)
        self.assertEqual(set(result["errors"][0]), {"row", "column", "message"})

    def test_formulas_errors_and_wrong_cell_types_are_invalid(self):
        for value in ("=1+1", "#VALUE!", 42, True):
            with self.subTest(value=value):
                self.write([valid_case(Title=value)])
                self.assert_column_error(self.validate(), "Title")

    def test_missing_duplicate_and_unknown_headers(self):
        for headers, column in (
            (HEADERS[1:], "Case ID"),
            (HEADERS + ["Title"], "Title"),
            (HEADERS + ["Unknown"], "Header 17"),
        ):
            with self.subTest(headers=headers):
                self.write(headers=headers)
                result = self.validate()
                self.assert_column_error(result, column, None)
                self.assertEqual(result["checked_rows"], 0)

    def test_missing_sheet_merged_cells_and_unheaded_data(self):
        for defect, column in (("sheet", "Test Cases"), ("merged", "Test Cases"),
                               ("unheaded", "Header 17")):
            with self.subTest(defect=defect):
                self.write()
                workbook = load_workbook(self.path)
                sheet = workbook["Test Cases"]
                if defect == "sheet":
                    sheet.title = "Other"
                elif defect == "merged":
                    sheet.merge_cells("A2:B2")
                else:
                    sheet["Q2"] = "unheaded"
                workbook.save(self.path)
                workbook.close()
                self.assert_column_error(self.validate(), column, None)

    def test_unreadable_missing_and_unsupported_workbooks(self):
        self.assert_column_error(self.validate(), "Workbook", None)
        self.path.write_text("not an xlsx", encoding="utf-8")
        self.assert_column_error(self.validate(), "Workbook", None)
        result = validate_test_catalog(self.path.with_suffix(".xlsm"), JIRA_URL, PROJECT)
        self.assert_column_error(result, "Workbook", None)

    def test_invalid_configuration_raises(self):
        for url, key in (
            ("http://jira.example", PROJECT), ("https://user:secret@jira.example", PROJECT),
            ("https://jira.example:bad", PROJECT), ("https://jira.example:0", PROJECT),
            ("", PROJECT), (JIRA_URL, ""), (JIRA_URL, "lower"), (JIRA_URL, "P-1"),
        ):
            with self.subTest(url=url, key=key):
                with self.assertRaises(ValueError):
                    validate_test_catalog(self.path, url, key)

    def test_cli_json_and_exit_codes(self):
        for valid in (True, False):
            self.write([valid_case()] if valid else [valid_case(Title="")])
            stdout = io.StringIO()
            with patch.object(sys, "argv", [
                "validate_test_catalog.py", str(self.path),
                "--jira-base-url", JIRA_URL, "--project-key", PROJECT,
            ]), contextlib.redirect_stdout(stdout):
                with self.assertRaises(SystemExit) as exit_result:
                    main()
            self.assertEqual(exit_result.exception.code, 0 if valid else 1)
            self.assertEqual(json.loads(stdout.getvalue())["outcome"], "valid" if valid else "invalid")

    def test_cli_configuration_failure_is_explicit(self):
        with patch.object(sys, "argv", [
            "validate_test_catalog.py", str(self.path),
            "--jira-base-url", "http://jira.example", "--project-key", PROJECT,
        ]), contextlib.redirect_stderr(io.StringIO()) as stderr:
            with self.assertRaises(SystemExit) as exit_result:
                main()
        self.assertEqual(exit_result.exception.code, 2)
        self.assertIn("HTTPS", stderr.getvalue())


if __name__ == "__main__":
    unittest.main()
