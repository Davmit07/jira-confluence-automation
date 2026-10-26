import contextlib
import io
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from create_shop_catalog import (
    DEFAULT_JIRA_URL, DEFAULT_PROJECT_KEY, SCENARIOS, create_shop_catalog, main,
)
from create_test_catalog import HEADERS
from validate_test_catalog import validate_test_catalog


class CreateShopCatalogTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name) / "shop.xlsx"

    def test_creates_valid_drafts_for_all_journeys_without_automation_claims(self):
        self.assertEqual(create_shop_catalog(self.path), self.path)
        result = validate_test_catalog(self.path, DEFAULT_JIRA_URL, DEFAULT_PROJECT_KEY)
        self.assertEqual(result, {
            "outcome": "valid", "checked_rows": 12, "error_count": 0, "errors": [],
        })
        workbook = load_workbook(self.path)
        try:
            rows = [
                dict(zip(HEADERS, row))
                for row in workbook["Test Cases"].iter_rows(min_row=2, values_only=True)
                if row[0]
            ]
            self.assertEqual(len(rows), len(SCENARIOS))
            self.assertEqual({row["Journey"] for row in rows}, {"Search", "Compare", "Configure", "Booking"})
            self.assertEqual(len({row["Case ID"] for row in rows}), 12)
            for row in rows:
                self.assertEqual(row["Review Status"], "Draft")
                self.assertEqual(row["Automation Status"], "Not Mapped")
                self.assertEqual(row["Last Result"], "Not Run")
                self.assertIsNone(row["Automation ID"])
                self.assertIsNone(row["Reviewer"])
                self.assertIsNone(row["Review Date"])
                self.assertTrue(row["Title"].startswith("[SYNTHETIC]"))
                self.assertIn("not verified", row["Preconditions"])
                if row["Journey"] == "Booking":
                    self.assertIn("Sandbox only", row["Preconditions"])
                    self.assertIn("simulated booking/payment", row["Preconditions"])
            self.assertEqual(len(workbook["Test Cases"].data_validations.dataValidation), 5)
            guidance = " ".join(str(row[1]) for row in workbook["Read Me"].values)
            self.assertIn("not fetched or verified", guidance)
            self.assertIn("fixed automation registry", guidance)
        finally:
            workbook.close()

    def test_refuses_overwrite_then_supports_explicit_force(self):
        self.path.write_bytes(b"existing")
        with self.assertRaises(FileExistsError):
            create_shop_catalog(self.path)
        self.assertEqual(self.path.read_bytes(), b"existing")
        create_shop_catalog(self.path, overwrite=True)
        self.assertEqual(validate_test_catalog(
            self.path, DEFAULT_JIRA_URL, DEFAULT_PROJECT_KEY
        )["outcome"], "valid")

    def test_custom_jira_configuration(self):
        create_shop_catalog(self.path, "https://jira.example:8443/projects/SHOP/", "SHOP")
        self.assertEqual(validate_test_catalog(
            self.path, "https://jira.example:8443", "SHOP"
        )["outcome"], "valid")

    def test_rejects_invalid_output_configuration_and_missing_directory(self):
        with self.assertRaises(ValueError):
            create_shop_catalog(self.path.with_suffix(".xls"))
        with self.assertRaises(ValueError):
            create_shop_catalog(self.path, "http://jira.example")
        with self.assertRaises(FileNotFoundError):
            create_shop_catalog(self.path.parent / "missing" / "shop.xlsx")
        self.assertFalse(self.path.exists())

    def test_failed_validation_does_not_publish_or_overwrite_and_cleans_temp_files(self):
        self.path.write_bytes(b"existing")
        with patch("create_shop_catalog.validate_test_catalog", return_value={
            "outcome": "invalid", "errors": [{"message": "invalid generated row"}],
        }):
            with self.assertRaisesRegex(ValueError, "failed validation"):
                create_shop_catalog(self.path, overwrite=True)
        self.assertEqual(self.path.read_bytes(), b"existing")
        self.assertEqual(list(self.path.parent.iterdir()), [self.path])

    def test_concurrent_destination_creation_cannot_be_overwritten(self):
        from os import link

        def competing_writer(source, destination):
            Path(destination).write_bytes(b"another user's catalog")
            link(source, destination)

        with patch("create_shop_catalog.os.link", side_effect=competing_writer):
            with self.assertRaises(FileExistsError):
                create_shop_catalog(self.path)
        self.assertEqual(self.path.read_bytes(), b"another user's catalog")
        self.assertEqual(list(self.path.parent.iterdir()), [self.path])

    def test_cli_generates_catalog_and_reports_errors(self):
        with patch.object(sys, "argv", ["create_shop_catalog.py", str(self.path)]), \
                contextlib.redirect_stdout(io.StringIO()) as stdout:
            main()
        self.assertIn("12 cases", stdout.getvalue())
        with patch.object(sys, "argv", ["create_shop_catalog.py", str(self.path)]), \
                contextlib.redirect_stderr(io.StringIO()) as stderr:
            with self.assertRaises(SystemExit) as exit_result:
                main()
        self.assertEqual(exit_result.exception.code, 2)
        self.assertIn("--force", stderr.getvalue())


if __name__ == "__main__":
    unittest.main()
