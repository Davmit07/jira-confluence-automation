import argparse
from pathlib import Path
from typing import Union

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter


DEFAULT_OUTPUT = Path("test-cases-template.xlsx")
MAX_CASE_ROWS = 1000
HEADERS = [
    "Case ID",
    "Jira Issue Key",
    "Jira Issue Link",
    "Journey",
    "Title",
    "Priority",
    "Preconditions",
    "Steps",
    "Expected Results",
    "Test Data Reference",
    "Review Status",
    "Reviewer",
    "Review Date",
    "Automation Status",
    "Automation ID",
    "Last Result",
]
LISTS = {
    "Journey": ["Search", "Compare", "Configure", "Booking"],
    "Priority": ["Critical", "High", "Medium", "Low"],
    "Review Status": ["Draft", "Needs Update", "Approved", "Rejected"],
    "Automation Status": ["Not Mapped", "Planned", "Implemented", "Retired"],
    "Last Result": ["Not Run", "Passed", "Failed", "Skipped"],
}
COLUMN_WIDTHS = [18, 20, 40, 16, 36, 18, 40, 55, 45, 30, 20, 24, 16, 22, 30, 16]


def create_test_catalog(
    output_path: Union[str, Path] = DEFAULT_OUTPUT, overwrite: bool = False
) -> Path:
    """Create a formatted Excel test-case catalog template."""
    output_path = Path(output_path)
    if output_path.exists() and not overwrite:
        raise FileExistsError(
            f"{output_path} already exists; choose another path or pass --force."
        )
    if not output_path.parent.is_dir():
        raise FileNotFoundError(f"Output directory does not exist: {output_path.parent}")

    workbook = Workbook()
    cases = workbook.active
    cases.title = "Test Cases"
    cases.append(HEADERS)
    cases.freeze_panes = "A2"
    cases.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{MAX_CASE_ROWS + 1}"
    cases.sheet_view.showGridLines = False
    cases.row_dimensions[1].height = 32

    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    for cell in cases[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for index, width in enumerate(COLUMN_WIDTHS, start=1):
        cases.column_dimensions[get_column_letter(index)].width = width

    for row in range(2, MAX_CASE_ROWS + 2):
        for column in range(1, len(HEADERS) + 1):
            cases.cell(row=row, column=column).alignment = Alignment(
                vertical="top", wrap_text=True
            )
        cases.cell(row=row, column=13).number_format = "yyyy-mm-dd"

    for header, options in LISTS.items():
        column = HEADERS.index(header) + 1
        validation = DataValidation(
            type="list", formula1=f'"{",".join(options)}"', allow_blank=True
        )
        validation.error = f"Select a value from the {header} list."
        validation.errorTitle = f"Invalid {header}"
        validation.showErrorMessage = True
        validation.prompt = f"Choose a {header.lower()}."
        validation.promptTitle = header
        validation.showInputMessage = True
        cases.add_data_validation(validation)
        validation.add(f"{get_column_letter(column)}2:{get_column_letter(column)}{MAX_CASE_ROWS + 1}")

    instructions = workbook.create_sheet("Read Me")
    instructions.sheet_view.showGridLines = False
    instructions.column_dimensions["A"].width = 28
    instructions.column_dimensions["B"].width = 110
    guidance = [
        ("Test Case Catalog", "Use one row per test case on the Test Cases sheet."),
        ("Case ID", "Assign a stable, unique ID and keep it unchanged when editing the case."),
        ("Jira traceability", "Enter the source issue key and its full URL for every case."),
        ("Journey / Priority", "Choose from the dropdowns so cases can be consistently filtered."),
        ("Steps", "Enter one numbered action per line in the same cell, for example: 1. Open search; 2. Enter criteria; 3. Submit."),
        ("Expected Results", "Describe observable outcomes for the case; use numbered lines when they correspond to steps."),
        ("Test Data Reference", "Use a safe named reference to test data; never put passwords, tokens, or real customer/payment data here."),
        ("Review Status", "New cases remain Draft until reviewed. Only QA-approved cases may be executed."),
        ("Review Date", "Enter the date the reviewer made the current decision."),
        ("Automation Status / ID", "Only mark a case Implemented when its Automation ID maps to implemented code. Workbook text is never executable."),
        ("Last Result", "Update after execution. Not Run is the initial state for a new case."),
        ("Booking safeguard", "Booking cases may only run against the approved Sandbox URL, regardless of other environment selection."),
        ("Sharing", "Avoid simultaneous edits to one workbook copy. Use the team's approved shared-storage and versioning process."),
    ]
    for row in guidance:
        instructions.append(row)
    for cell in instructions["A"]:
        cell.font = Font(bold=True, color="1F1F1F")
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    for row in instructions.iter_rows(min_row=1, max_col=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    instructions.freeze_panes = "A2"

    workbook.properties.title = "Vehicle Shop Test Cases Template"
    workbook.properties.subject = "Jira-traceable, reviewable test-case catalog"
    workbook.save(output_path)
    return output_path


def main() -> None:
    parser = argparse.ArgumentParser(description="Create the Vehicle Shop Excel test-case template.")
    parser.add_argument(
        "output",
        nargs="?",
        type=Path,
        default=DEFAULT_OUTPUT,
        help=f"output workbook path (default: {DEFAULT_OUTPUT})",
    )
    parser.add_argument(
        "--force",
        action="store_true",
        help="overwrite the output workbook if it already exists",
    )
    arguments = parser.parse_args()
    try:
        output_path = create_test_catalog(arguments.output, overwrite=arguments.force)
    except (FileExistsError, FileNotFoundError) as error:
        parser.error(str(error))
    print(f"Created test catalog template: {output_path}")


if __name__ == "__main__":
    main()
