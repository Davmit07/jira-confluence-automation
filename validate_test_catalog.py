"""Read-only, offline validation for the Vehicle Shop Excel catalog."""

import argparse
import json
import re
from datetime import date, datetime
from pathlib import Path
from typing import List, Optional, Tuple, TypedDict, Union
from urllib.parse import urlsplit
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from create_test_catalog import HEADERS, LISTS


REQUIRED_FIELDS = (
    "Case ID", "Jira Issue Key", "Jira Issue Link", "Journey", "Title",
    "Priority", "Preconditions", "Steps", "Expected Results",
    "Review Status", "Automation Status", "Last Result",
)


class ValidationError(TypedDict):
    row: Optional[int]
    column: str
    message: str


class ValidationResult(TypedDict):
    outcome: str
    checked_rows: int
    error_count: int
    errors: List[ValidationError]


def jira_origin(url: str) -> Tuple[str, str, int]:
    """Return an HTTPS origin; accept a Jira project URL as configuration."""
    if (
        not isinstance(url, str) or not url
        or any(character.isspace() for character in url) or "\\" in url
    ):
        raise ValueError("Jira base URL must be an absolute HTTPS URL without whitespace.")
    try:
        parsed = urlsplit(url)
        port = parsed.port
    except ValueError:
        raise ValueError("Jira base URL is malformed.") from None
    if (
        parsed.scheme != "https" or not parsed.hostname
        or parsed.username is not None or parsed.password is not None
        or parsed.query or parsed.fragment or port == 0
    ):
        raise ValueError("Jira base URL must use HTTPS without credentials, query, or fragment.")
    return parsed.scheme, parsed.hostname, port if port is not None else 443


def validate_configuration(jira_base_url: str, project_key: str) -> None:
    jira_origin(jira_base_url)
    if not isinstance(project_key, str) or not re.fullmatch(r"[A-Z][A-Z0-9_]*", project_key):
        raise ValueError("Project key must start with A-Z and contain only A-Z, 0-9, or _.")


def _valid_steps(value: str) -> bool:
    lines = value.splitlines()
    if not lines:
        return False
    for expected_number, line in enumerate(lines, start=1):
        match = re.fullmatch(r"\s*([1-9][0-9]*)[.)]\s+(\S.*)", line)
        if not match or match.group(1) != str(expected_number):
            return False
    return True


def _valid_review_date(value) -> bool:
    if isinstance(value, (datetime, date)):
        return True
    if not isinstance(value, str) or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
        return False
    try:
        date.fromisoformat(value)
    except ValueError:
        return False
    return True


def validate_test_catalog(
    catalog_path: Union[str, Path], jira_base_url: str, project_key: str
) -> ValidationResult:
    """Return outcome/counts/errors; never change the workbook or contact Jira.

    Errors use Excel row numbers and exact schema headers. Catalog-level errors
    have row=None and identify the sheet/header in column. Invalid configuration
    is a caller error and raises ValueError, rather than returning a valid result.
    """
    validate_configuration(jira_base_url, project_key)
    errors: List[ValidationError] = []
    checked_rows = 0

    def error(row: Optional[int], column: str, message: str) -> None:
        errors.append({"row": row, "column": column, "message": message})

    def result() -> ValidationResult:
        return {
            "outcome": "invalid" if errors else "valid",
            "checked_rows": checked_rows,
            "error_count": len(errors),
            "errors": errors,
        }

    catalog_path = Path(catalog_path)
    if catalog_path.suffix.lower() != ".xlsx":
        error(None, "Workbook", "Unsupported format; provide an .xlsx catalog.")
        return result()
    try:
        workbook = load_workbook(catalog_path, data_only=False, keep_links=False)
    except (OSError, BadZipFile, InvalidFileException, ParseError, ValueError, KeyError):
        error(None, "Workbook", "Cannot read workbook; provide a readable, unencrypted .xlsx catalog.")
        return result()

    try:
        if "Test Cases" not in workbook.sheetnames:
            error(None, "Test Cases", "Required worksheet is missing.")
            return result()
        cases = workbook["Test Cases"]
        if not hasattr(cases, "iter_rows"):
            error(None, "Test Cases", "Test Cases must be a cell worksheet.")
            return result()
        if cases.merged_cells.ranges:
            error(None, "Test Cases", "Merged cells are unsupported; use one cell per field.")
        header_values = [cell.value for cell in cases[1]]
        for header in HEADERS:
            count = header_values.count(header)
            if count == 0:
                error(None, header, "Required header is missing.")
            elif count > 1:
                error(None, header, "Duplicate header; keep exactly one column with this name.")
        for index, header in enumerate(header_values):
            if header is not None and header not in HEADERS:
                error(None, f"Header {index + 1}", "Unsupported header; use the exact template headers.")
            elif header is None and any(
                cell.value is not None
                for (cell,) in cases.iter_rows(min_row=2, min_col=index + 1, max_col=index + 1)
            ):
                error(None, f"Header {index + 1}", "Data column has no header.")
        if errors:
            return result()

        columns = {header: header_values.index(header) for header in HEADERS}
        seen_ids = {}
        issue_pattern = re.compile(re.escape(project_key) + r"-[1-9][0-9]*")
        configured_origin = jira_origin(jira_base_url)
        for row_number, cells in enumerate(cases.iter_rows(min_row=2), start=2):
            if all(cell.value is None or (
                isinstance(cell.value, str) and not cell.value.strip()
            ) for cell in cells):
                continue
            checked_rows += 1
            values = {header: cells[index].value for header, index in columns.items()}
            text = {
                header: value if isinstance(value, str) else ""
                for header, value in values.items()
            }
            for header in HEADERS:
                cell = cells[columns[header]]
                if cell.data_type in ("f", "e"):
                    error(row_number, header, "Formulas and Excel error cells are unsupported; enter a literal value.")
                elif values[header] is not None and header != "Review Date" and not isinstance(values[header], str):
                    error(row_number, header, "Enter text, not a numeric or other cell value.")
            for header in REQUIRED_FIELDS:
                if not text[header].strip():
                    error(row_number, header, "Required field must contain text.")
            case_id = text["Case ID"].strip().casefold()
            if case_id:
                if case_id in seen_ids:
                    error(row_number, "Case ID", f"Duplicate case ID; first used on row {seen_ids[case_id]}.")
                else:
                    seen_ids[case_id] = row_number
            for header, options in LISTS.items():
                if text[header].strip() and text[header] not in options:
                    error(row_number, header, "Choose exactly one of: " + ", ".join(options) + ".")
            issue_key = text["Jira Issue Key"]
            if issue_key.strip() and not issue_pattern.fullmatch(issue_key):
                error(row_number, "Jira Issue Key", f"Use {project_key}- followed by a positive issue number.")
            link = text["Jira Issue Link"]
            if link.strip():
                try:
                    link_origin = jira_origin(link)
                    parsed_link = urlsplit(link)
                    valid_link = (
                        link_origin == configured_origin
                        and bool(issue_pattern.fullmatch(issue_key))
                        and parsed_link.path == f"/browse/{issue_key}"
                    )
                except ValueError:
                    valid_link = False
                if not valid_link:
                    error(row_number, "Jira Issue Link", "Use an HTTPS /browse/ URL on the configured Jira origin matching Jira Issue Key.")
            if text["Steps"].strip() and not _valid_steps(text["Steps"]):
                error(row_number, "Steps", "Use one numbered action per line, starting at 1 with no gaps or duplicates.")
            if text["Review Status"] == "Approved":
                if not text["Reviewer"].strip():
                    error(row_number, "Reviewer", "Approved cases require a reviewer.")
                if values["Review Date"] is None or values["Review Date"] == "":
                    error(row_number, "Review Date", "Approved cases require a review date.")
            if values["Review Date"] is not None and values["Review Date"] != "":
                if not _valid_review_date(values["Review Date"]):
                    error(row_number, "Review Date", "Use an Excel date or a valid YYYY-MM-DD date.")
            if text["Automation Status"] == "Implemented" and not text["Automation ID"].strip():
                error(row_number, "Automation ID", "Implemented cases require an automation ID; registry checks belong to the runner.")
        return result()
    finally:
        workbook.close()


def main() -> None:
    parser = argparse.ArgumentParser(description="Validate a Vehicle Shop catalog offline; never modify it.")
    parser.add_argument("catalog", type=Path)
    parser.add_argument("--jira-base-url", required=True, help="HTTPS Jira origin or project URL.")
    parser.add_argument("--project-key", required=True)
    arguments = parser.parse_args()
    try:
        validation = validate_test_catalog(
            arguments.catalog, arguments.jira_base_url, arguments.project_key
        )
    except ValueError as error:
        parser.error(str(error))
    print(json.dumps(validation, indent=2))
    raise SystemExit(0 if validation["outcome"] == "valid" else 1)


if __name__ == "__main__":
    main()
